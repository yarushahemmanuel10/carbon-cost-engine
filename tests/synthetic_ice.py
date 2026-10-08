"""Builds a small, MADE-UP workbook laid out like the ICE summary sheet.

No number in here comes from ICE: the values are arbitrary and are not realistic carbon factors. The layout copies the
real file's quirks so the parser is tested without the licensed data:
heading levels shown only by formatting, header variants, a title-style
header, multi-value columns, a headerless table, a withdrawn section, the
interactive block, the timber pair, and the duplicate cases.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

THEME = 9
COLS = {c: i + 1 for i, c in enumerate("ABCDEFGHIJKLMNOPQRSTUVWXYZ")}
NAME = "Name Exactly as in ICE DB"


def _fill(tint):
    return PatternFill(fill_type="solid", fgColor=Color(theme=THEME, tint=tint))


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.n = 0
        self.rows = {}          # label -> row number, for assertions

    def blank(self, count=1):
        self.n += count

    def _write(self, cells, bold=True, size=14.0, fill=None, tag=None):
        self.n += 1
        for col, value in cells.items():
            self.ws.cell(row=self.n, column=COLS[col], value=value)
        f = self.ws.cell(row=self.n, column=COLS["F"])
        f.font = Font(bold=bold, size=size)
        if fill is not None:
            f.fill = _fill(fill)
        if tag:
            self.rows[tag] = self.n
        return self.n

    def category(self, text):
        self.blank()
        self._write({"F": text}, bold=False, size=20.0, fill=0.4)
        self.blank()

    def section(self, text):
        self._write({"F": text}, fill=0.6)
        self.blank()

    def sub(self, text):
        self._write({"F": text}, size=16.0)
        self.blank()

    def band(self, text, **cells):
        return self._write({"F": text, **cells}, fill=0.8)

    def note(self, text, bold=False):
        self._write({"F": text}, bold=bold)
        self.blank()

    def version(self, text):
        self._write({"F": f"Version: {text}"})
        self.blank()

    def header(self, **cells):
        return self._write({"F": "Materials", **cells}, fill=0.8)

    def row(self, label, tag=None, **cells):
        return self._write({"F": label, **cells}, tag=tag or label)


def build(path, select_unit="kg", include_selector=True, version_text="Test (V5.0) - made up"):
    wb = Workbook()
    wb.active.title = "Cover"
    s = Sheet(wb.create_sheet("ICE Summary"))
    KG = "Embodied Carbon -  kgCO2e/kg"
    std = dict(G=KG, J="Comments", K="DQI Score", L="DQI Version", T=NAME)

    s.blank(5)
    s._write({"F": "TEST SUMMARY - made-up data"}, size=24.0)
    s.row("ICE DB Version:", G=version_text)
    s.row("Note on scores", G="This text sits before the first category")

    # ---------------- simple per-kg table; a text value; a blank value; a repeat (rule B)
    s.category("Aggregates")
    s.header(**std)
    s.row("gravel", G=0.5, J="a comment", K=0.6, L="V1.0", T="Aggregate, gravel")
    s.row("sand", G=0.4, K="N/A", L="N/A", T="Aggregate, sand")
    s.row("removed stuff", G="Data has been removed from this dataset.", T="Aggregate, removed")
    s.row("no value given", T="Aggregate, blank")

    # ---------------- per tonne
    s.category("Tarmac")
    s.band("Tarmac, mixed at plant")
    s.blank()
    s.note("The boundaries of this data are cradle to mixed tarmac. It includes mixing energy.")
    s.header(E="Embodied Carbon -  kgCO2e per tonne", G="Embodied Carbon -  kgCO2e per tonne", K="DQI Score", L="DQI Version", T=NAME)
    s.row("Tarmac, 5% binder", G=80.0, K=0.7, L="V1.0", T="Tarmac, 5% binder")
    s.blank()
    # per unit, with a 'Units' column; the second row has an Excel error there
    s.band("Road surfaces")
    s.blank()
    s.header(E="Embodied Carbon -  kgCO2e per 1 m2, per 100mm depth", G="Embodied Carbon -  kgCO2e per unit", I="Units",
             J="Comments", K="DQI Score", L="DQI Version", T=NAME)
    s.row("Road, 5% binder", G=21.0, I="1 m2, per 100mm depth", T="Road, 5% binder")
    s.row("Road, 6% binder", G=23.0, I="#N/A", T="Road, 6% binder")

    # ---------------- bricks: per-kg trap, and "1" + "m2" declared units
    s.category("Bricks")
    s.header(**std)
    s.row("General brick", G=0.9, T="Clay, Brick")
    s.row("A Single Brick", G=1.9, J="a brick weighs about 2 kg", T="Clay brick, one brick, 215 x 65 x 102.5mm")
    s.row("A single slab", G=2.9, T="Slab, general")            # only the label says 'one item'
    s.blank()
    s.section("Brick walls - per m2 of wall")
    s.sub("Single skin walls")
    s.header(E="Embodied Carbon -  kgCO2e per m2", G="Declared Units", H="EC -  kgCO2e per unit", I="Mass per declared unit - kg",
             J="Comments", K="DQI Score", L="DQI Version", T=NAME)
    s.row("with mortar 1:3", tag="brick wall 1:3", G=1, H=60.0, I="m2", T="wall, brick, single skin, mortar 1:3")
    s.row("with mortar 1:4", tag="brick wall 1:4", G=2, H=61.0, I="m2", T="wall, brick, single skin, mortar 1:4")

    # ---------------- cement: sub-section, bands inside a table, and a HEADERLESS table
    s.category("Cement and Mortar")
    s.sub("Cement")
    s.header(**std)
    s.row("General cement", G=3.0, T="Cement, general")
    s.band("CEM II", T=NAME)
    s.row("CEM II/A", G=2.5, T="Cement CEM II/A")
    s.blank()
    s.sub("Mortar")
    s.band("Mortar - CEM I", K="DQI Score", L="DQI Version", T=NAME)       # no 'Materials' header, no unit
    s.row("Mortar (1:3)", G=0.6, K=0.7, L="V1.0", T="Mortar (1:3)")
    s.row("Mortar (1:4)", G=0.5, K=0.7, L="V1.0", T="Mortar (1:4)")

    # ---------------- withdrawn section: header with no value column
    s.category("Carpet")
    s.version("ICE V 2.0 2011")
    s.header(G="Comments", T=NAME)
    s.row("Felt underlay")
    s.row("Nylon carpet", G="see the note above")

    # ---------------- concrete
    s.category("Concrete")
    s.section("In-Situ Concrete")
    if include_selector:
        s.row("SELECT declared unit", G=select_unit)
        s.blank()
    s.section("By Strength Class")
    s.note("Important Note on Data By Strength Class: \n\nUse the strength class.", bold=True)
    s.band("In-Situ Concrete", G="Per kg of concrete", J="Comments", K="DQI Score", L="DQI Version")   # title row
    s.header(G="Embodied Carbon -  kgCO2e per kg", T=NAME, U=NAME, V=NAME)
    s.row("General", tag="concrete general", G=0.31, J="avoid general values", K=0.71, L="V1.0", T="Concrete, average mix")
    s.row("32/40 MPa", tag="avg 32/40", G=0.33, K=0.7, L="V1.0", T="Concrete 32/40 average")
    s.band("With CEM I", T=NAME)
    s.row("32/40 MPa", tag="cem1 32/40", G=0.34, K=0.7, L="V1.0", T="Concrete 32/40 with CEM I")
    s.band("% Cement Replacement - Slag", G=0.25, H=0.5, J="Slag", K="DQI Score\n(25%)", L="DQI Score\n(50%)",
           N="DQI Version", T=0.25, U=0.5)
    s.row("32/40 MPa", tag="slag 32/40", G=0.27, H=0.21, J="estimated", K=0.69, L=0.68, N="V1.0",
          T="Concrete 32/40 with 25% slag", U="Concrete 32/40 with 50% slag")
    s.blank()
    s.section("By Volume and Cementitious Content")
    s.band("In-Situ Concrete", E="Per kg of concrete", G="Embodied Carbon -  kgCO2e per kg", J="Comments", K="DQI Score",
           L="DQI Version", T=NAME)                                          # title-style header
    s.band("Nominal proportions - with average cement", T=NAME)
    s.row("1:2.5:5", tag="mix wrong label", G=0.44, T="1:3:6 mix with average cement")     # rule A
    s.row("1:3:6", tag="mix right label", G=0.44, T="1:3:6 mix with average cement")
    s.band("By cement content with CEM I", T=NAME)
    s.row("100 kg CEM I per m3", tag="cc 100", G=0.15, T="Concrete - CEM I - 100 kg")
    s.row(120, tag="cc 120", G=0.16, T="Concrete - CEM I - 120 kg")           # numeric label
    s.band("By cement content with slag")
    s.band("% Cement Replacement - Slag", G=0, H=0.5, J="Slag", K="DQI Score\n(0%)", L="DQI Score\n(50%)",
           N="DQI Version", T=0, U=0.5)
    s.row(120, tag="cc slag 120", G=0.16, H=0.12, K=0.7, L=0.66, N="V1.0",
          T="Concrete - CEM I - 120 kg", U="Concrete - 50% slag - 120 kg")    # G column repeats 'cc 120' (rule B)
    s.blank()
    s.section("Precast Concrete")
    s.sub("Precast concrete model")
    s.note("Here you can model any type of concrete as a precast concrete.")
    s.row("SELECT concrete type (scroll up on drop down list)", G="Results for precast concrete of selected type")
    s._write({"G": "Declared Units", "H": "Embodied Carbon -  kgCO2e per unit", "I": "Mass per declared unit - kg"})
    s.row("Concrete - model row", G="m3", H=999.0, I=1234)
    s.blank()
    s.note("NOTE: The below data has multiple units. Check each datapoint for the units", bold=True)
    s.section("Precast concrete products")
    s.header(G="Declared Units", H="Embodied Carbon -  kgCO2e per unit", I="Mass per declared unit - kg", J="Comments",
             K="DQI Score", L="DQI Version", T=NAME)
    s.row("precast pipe", G="m", H=12.0, I=300, T="precast pipe, per m")
    s.row("precast beam", G="kg", H=0.55, I=1, T="precast beam, per kg")
    s.row("As above but with recycled rebar", tag="as above 1", G="kg", H=0.50, I=1, T="precast beam, recycled rebar")
    s.row("precast floor", G="m2", H=77.0, I=333, T="precast floor, per m2")
    s.row("As above but with recycled rebar", tag="as above 2", G="m2", H=70.0, I=333, T="precast floor, recycled rebar")   # rule C
    s.row("concrete block 100mm", G="each", H=3.5, I=20.5, T="block, 100mm, each")

    # ---------------- glass: per m2 header
    s.category("Glass")
    s.header(**std)
    s.row("Glass, general, per kg", G=5.5, T="Glass, general, per kg")
    s.blank()
    s.section("Glazing by thickness - Per m2")
    s.band("Double glazing - no frame", E="Embodied Carbon -  kgCO2e per m2")
    s.blank()
    s.header(G="Embodied Carbon - kgCO2e/m2", J="Comments", K="DQI Score", L="DQI Version", T=NAME)
    s.row("8 mm of glass", G=88.0, T="Double, 8 mm of glass")

    # ---------------- steel: header without a name label; the same datapoint shown twice (rule B)
    s.category("Steel")
    s.band("World data - see the profile", T=NAME)
    s.blank()
    s.header(G=KG, K="DQI Score", L="DQI Version")
    s.row("Steel, Rebar", tag="rebar world", G=9.9, K=0.8, L="V1.0", T="Steel, Rebar")
    s.note("The data above has been extracted from a test database, see https://example.org/steel", bold=True)
    s.band("Regional data")
    s.blank()
    s.header(**std)
    s.row("Steel, Rebar", tag="rebar repeat", G=9.9, K=0.8, L="V1.0", T="Steel, Rebar")
    s.row("Steel rebar, recycled", G=4.4, K=0.6, L="V1.0", T="Steel rebar, recycled")

    # ---------------- timber: excluding / including carbon storage
    s.category("Timber")
    tim = dict(G=KG, I="Of which carbon storage - kg CO2e/kg", J="Comments", K="DQI Score", L="DQI Version", T=NAME)
    s.section("Data excluding carbon storage")
    s.note("The data in this section excludes the carbon stored in the timber.")
    s.header(**tim)
    s.row("Timber, CLT - No Carbon Storage", G=0.60, I=0, K=0.8, L="V1.0", T="Timber, CLT - No Carbon Storage")
    s.row("Timber, Glulam - No Carbon Storage", G=0.65, I=0, T="Timber, Glulam - No Carbon Storage")
    s.row("Timber, Odd - No Carbon Storage", G=0.80, I=0, T="Timber, Odd - No Carbon Storage")
    s.blank(2)
    s.section("Data including carbon storage")
    s.header(**tim)
    s.row("Timber, CLT - Including Carbon Storage", G=-1.40, I=-2.00, T="Timber, CLT - Including Carbon Storage")
    s.row("Timber, Glulam - Including Carbon Storage", G=-1.25, I=-1.90, T="Timber, Glulam - Including Carbon Storage")
    s.row("Timber, Odd - Including Carbon Storage", G=-0.90, I=-1.00, T="Timber, Odd - Including Carbon Storage")     # 0.8 - 1.0 is not -0.9
    s.row("Timber, Orphan - Including Carbon Storage", G=-1.00, I=-1.30, T="Timber, Orphan - Including Carbon Storage")
    s.blank()
    s._write({"F": "<<< End of Summary >>>"}, size=18.0)

    # ---------------- a profile sheet with a density table (made-up values)
    p = wb.create_sheet("Timber")
    p.cell(row=3, column=3, value="Material Properties")                   # a navigation link, not the table
    p.cell(row=10, column=6, value="Material Properties (CIBSE Data)")
    for col, text in (("F", "Material"), ("H", "Condition"), ("I", "Thermal conductivity"), ("J", "Density (kg m-3)")):
        p.cell(row=11, column=COLS[col], value=text)
    data = [("softwood", None, 500), (None, "Dry", 600), ("balsa", None, 150), ("Wood derivatives:", None, None),
            (None, None, 700), ("plywood", "Moist", 550), (None, 0, 1200)]
    for i, (mat, cond, dens) in enumerate(data, start=12):
        p.cell(row=i, column=2, value=i - 11)
        if mat is not None:
            p.cell(row=i, column=COLS["F"], value=mat)
        if cond is not None:
            p.cell(row=i, column=COLS["H"], value=cond)
        if dens is not None:
            p.cell(row=i, column=COLS["J"], value=dens)
    for i in range(19, 30):
        p.cell(row=i, column=2, value=i - 11)                              # numbered but empty rows, as in the real file
    q = wb.create_sheet("Plaster")
    q.cell(row=5, column=6, value="Material Properties (CIBSE Data)")
    q.cell(row=6, column=6, value="Material")
    q.cell(row=6, column=10, value="Density (kg m-3)")
    q.cell(row=7, column=6, value="board")
    q.cell(row=7, column=10, value=800)
    wb.save(path)
    return s.rows

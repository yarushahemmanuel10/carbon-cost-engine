"""Ingest the ICE Database (Educational V5.0) summary sheet.

    python -m carbon_cost.ingest.ice                # extract, parse, write staging files, load
    python -m carbon_cost.ingest.ice --no-load      # everything except writing to the database

The ICE file is licensed for educational use only and may not be
redistributed. This script reads the copy YOU downloaded into data/raw/.
Everything it writes goes to data/staging/ and db/, which are git-ignored.

How the "ICE Summary" sheet is laid out (this is what the parser relies on)
---------------------------------------------------------------------------
* It is a stack of small tables. A table starts with a HEADER row, where
  column F is exactly "Materials". Data rows follow with no gap. The first
  row with an empty column F ends the table.
* The unit is only written in the header text ("kgCO2e/kg", "kgCO2e per
  tonne", "kgCO2e/m2", "kgCO2e per unit"), so the layout is read from each
  header row. Nothing is tied to a row number.
* Rows with text in F and nothing in G..I are headings or notes. Their
  level is not written anywhere; the only signal is cell formatting:
      font size 20 or more                  -> category       ("Concrete")
      fill tint about 0.6                   -> section        ("By Strength Class")
      bold, no fill, short                  -> sub-section    ("Cement")
      fill tint about 0.8 (same as header)  -> band           ("With CEM I")
  Anything else is a note and is ignored (but listed in ice_structure.csv).
* A header row normally has "Materials" in column F. One concrete table
  instead has a title there and the unit text in column G; it is treated
  as a header too. A band-coloured title row directly above a header can
  carry the "Comments" / "DQI" labels that the header itself leaves out.
* Some concrete tables give several values per row (for example 25%, 50%
  and 70% GGBS in columns G, H, I). They are introduced by a band row
  whose G..I cells hold the percentages. Each column becomes its own
  material.
"""
from __future__ import annotations

import argparse
import csv
import logging
import re
from collections import Counter
from dataclasses import dataclass, field, asdict
from pathlib import Path

from carbon_cost import db

log = logging.getLogger("carbon_cost.ingest.ice")

# ---------------------------------------------------------------- constants
PROJECT_ROOT = db.PROJECT_ROOT
DEFAULT_XLSX = PROJECT_ROOT / "data" / "raw" / "ICE_DB_Educational_V5.0_-_Jun_2026.xlsx"
STAGING_DIR = PROJECT_ROOT / "data" / "staging"

SUMMARY_SHEET = "ICE Summary"
N_COLS = 34                                        # columns A..AH
COL_E, COL_F, COL_G, COL_H, COL_I = 4, 5, 6, 7, 8

# Convention agreed for this project: every ICE V5.0 factor is valid from
# the month the database was released. See README, "Data sources".
ICE_VALID_FROM = "2026-06-01"
# The sheet states its own version next to the label "ICE DB Version:".
# Factors are stored as V5.0, so a file that says anything else is refused.
EXPECTED_SHEET_VERSION = "V5.0"
SHEET_VERSION_LABEL_RE = re.compile(r"^\s*ice db version\s*:?\s*$", re.I)
LIFECYCLE_MODULE = "A1-A3"                         # ICE is cradle-to-gate

ICE_SOURCE = {
    "name": "ICE Database Educational",
    "publisher": "Circular Ecology Ltd",
    "version": "V5.0 (Jun 2026)",
    "licence": "Educational use only. May not be redistributed in whole or in part.",
    "attribution_text": "ICE Database Educational V5.0, Circular Ecology Ltd",
    "url": "https://circularecology.com/embodied-carbon-footprint-database.html",
    "redistributable": 0,
}

# The sheet formats "In-Situ Concrete" and "By Strength Class" identically,
# although the second sits inside the first. Formatting cannot tell them
# apart, so the few headings that GROUP other sections are named here.
# Key = category, value = headings that act as a group inside it.
# If ICE renames them, the parser logs a warning and the path simply has
# one level fewer; no value is affected.
GROUP_HEADINGS = {
    "Concrete": ("In-Situ Concrete", "Precast Concrete"),
}

# The interactive calculator block. Its numbers depend on a drop-down
# selection, so nothing in it is a fixed factor.
PRECAST_MODEL_RE = re.compile(r"precast concrete model", re.I)

# In-situ concrete values change with this drop-down. It must read "kg".
DECLARED_UNIT_SELECTOR_RE = re.compile(r"^\s*select declared unit\s*$", re.I)

# A row such as "A Single Brick" that sits in a per-kg table. Its value is per
# item, not per kg, so loading it as per kg would be wrong by the item's mass.
#   - If the ICE DB name itself says it is one item ("Clay brick, one brick,
#     215 x 65 x 102.5mm"), the source states the unit: the row is stored per
#     unit, that name is the declared unit, and the conflict is noted in comments.
#   - If only the row label says so, the unit is not stated: the row is rejected.
SINGLE_ITEM_RE = re.compile(r"\b(a single|one)\s+(brick|block|tile|slab|unit|item|sheet|panel|pipe)\b", re.I)

TIMBER_SUFFIX_RE = re.compile(
    r"\s*-\s*(no|including|excluding|includes|excludes|without|with)\s+carbon\s+storage\s*$", re.I
)

# Reasons a row is rejected (short codes are used for the counts).
R_BLANK = "value is blank (data withdrawn or not given)"
R_TEXT = "value is text, not a number"
R_NO_VALUE_COL = "table has no value column (withdrawn ICE V2.0 section)"
R_PRECAST = "inside the interactive 'Precast concrete model' block (values depend on a drop-down)"
R_NO_HEADER = "numeric value but no header row above it, so the unit is not stated"
R_UNIT = "unit could not be read"
R_NEGATIVE = "negative value"
R_TIMBER_UNPAIRED = "timber 'including carbon storage' row with no matching 'excluding' row"
R_TIMBER_NO_HEADING = "timber table without an excluding/including heading"
R_DUPLICATE = "repeat of another row (same ICE DB name, same value)"
R_UNIT_CONFLICT = ("row label describes one single item but its ICE DB name does not, and the table header "
                   "says per kg (unit conflict in the source)")
UNIT_FROM_NAME_NOTE = ("[Note added on ingestion: the ICE table header says kgCO2e per kg, but the ICE DB name "
                       "describes one single item. Stored as kgCO2e per unit; the unit is the item named in "
                       "declared_unit.]")
R_LABEL_MISMATCH = "row label disagrees with its ICE DB name, and the value repeats another row"


class IngestError(RuntimeError):
    """Raised when the file does not look the way the parser needs it to."""


# ---------------------------------------------------------------- data classes
@dataclass
class SheetRow:
    row: int                    # row number in the sheet (1-based)
    cells: list                 # values of columns A..AH
    bold: bool = False          # formatting of the cell in column F
    size: float = 0.0
    fill: str = ""              # e.g. 'theme9:0.60', 'rgb:FFFF0000' or ''

    @property
    def f(self):
        """Column F as text. Some labels are plain numbers (cement content '120'), so numbers count."""
        v = self.cells[COL_F]
        if isinstance(v, str):
            return v.strip()
        if isinstance(v, bool) or v is None:
            return None
        if isinstance(v, (int, float)):
            return f"{v:g}"
        return str(v)


@dataclass
class Layout:
    header_row: int
    signature: tuple
    value_col: int | None = None
    unit_kind: str | None = None        # 'kg', 'tonne', 'm2', 'unit'
    unit_text: str = ""
    declared_col: int | None = None
    mass_col: int | None = None
    unit_desc_col: int | None = None
    storage_col: int | None = None
    comments_col: int | None = None
    dqi_col: int | None = None
    dqi_version_col: int | None = None
    name_col: int | None = None
    table_declared_unit: str | None = None   # from the text in column E of the header row


@dataclass
class ValueColumn:
    """One value column of a multi-value table (e.g. the '50% GGBS' column)."""
    col: int
    share: float                # 0.5 for 50%
    dqi_col: int | None
    name_col: int | None

    @property
    def label(self) -> str:
        return f"{self.share * 100:g}%"


@dataclass
class ParsedFactor:
    row: int
    category: str
    subcategory_path: str
    label: str
    canonical_name: str
    name_in_source: str | None
    value: float                # in `unit`
    unit: str                   # kgCO2e/kg | kgCO2e/m2 | kgCO2e/unit
    declared_unit: str | None
    mass_per_item_kg: float | None
    source_value: float         # exactly as in the sheet
    source_unit_text: str       # header text the unit was read from
    comments: str | None
    dqi_score: float | None
    dqi_version: str | None
    section_version: str | None
    biogenic_kgco2e_per_kg: float | None = None
    biogenic_row: int | None = None
    note: str = ""


@dataclass
class Rejected:
    row: int
    path: str
    name: str
    raw_value: str
    reason: str
    detail: str = ""


@dataclass
class ParseResult:
    factors: list = field(default_factory=list)
    rejected: list = field(default_factory=list)
    structure: list = field(default_factory=list)      # (row, kind, text)
    header_variants: Counter = field(default_factory=Counter)
    duplicates: list = field(default_factory=list)     # dicts describing each resolution
    timber_pairs: int = 0
    warnings: list = field(default_factory=list)
    declared_unit_selector: str | None = None
    sheet_version: str | None = None      # text next to 'ICE DB Version:' at the top of the sheet
    unit_from_name: list = field(default_factory=list)   # rows stored per unit because the ICE DB name says so


# ---------------------------------------------------------------- step 1: read the sheet
def _fill_key(cell) -> str:
    try:
        fill = cell.fill
        if fill is None or fill.fill_type is None:
            return ""
        colour = fill.fgColor
        if colour.type == "theme":
            return f"theme{colour.theme}:{round(colour.tint or 0.0, 2):.2f}"
        if colour.type == "rgb":
            return f"rgb:{colour.rgb}"
        return str(colour.type)
    except Exception:           # empty cells carry no style
        return ""


def read_summary_sheet(xlsx_path: Path) -> list[SheetRow]:
    """Read every non-empty row of the 'ICE Summary' sheet (values + column F formatting)."""
    import warnings

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise IngestError(
            f"ICE file not found: {xlsx_path}\n"
            "Download 'ICE Database Educational V5.0' from "
            f"{ICE_SOURCE['url']} and put it in data/raw/."
        )
    # read_only: some sheets report about a million rows. data_only: cached values, not formulas.
    wb = _open_workbook(xlsx_path)
    try:
        if SUMMARY_SHEET not in wb.sheetnames:
            raise IngestError(f"Sheet '{SUMMARY_SHEET}' not found. Sheets present: {wb.sheetnames}")
        ws = wb[SUMMARY_SHEET]
        rows: list[SheetRow] = []
        warnings.simplefilter("ignore")            # openpyxl warns about conditional formatting it cannot read
        for number, cells in enumerate(ws.iter_rows(min_row=1, max_col=N_COLS), start=1):
            values = [c.value for c in cells]
            values += [None] * (N_COLS - len(values))
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            f_cell = cells[COL_F] if len(cells) > COL_F else None
            bold, size, fill = False, 0.0, ""
            if f_cell is not None and f_cell.value is not None:
                try:
                    bold = bool(f_cell.font.b)
                    size = float(f_cell.font.sz or 0)
                except Exception:
                    pass
                fill = _fill_key(f_cell)
            rows.append(SheetRow(number, values, bold, size, fill))
    finally:
        wb.close()
        warnings.resetwarnings()
    return rows


def _open_workbook(xlsx_path: Path):
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)


def write_raw_csv(rows: list[SheetRow], path: Path) -> None:
    letters = [_col_letter(i) for i in range(N_COLS)]
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["row", "f_bold", "f_size", "f_fill", *letters])
        for r in rows:
            w.writerow([r.row, int(r.bold), r.size, r.fill, *["" if v is None else v for v in r.cells]])


def _col_letter(index: int) -> str:
    return chr(65 + index) if index < 26 else "A" + chr(65 + index - 26)


# ---------------------------------------------------------------- step 2: parse
def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _text(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def _is_excel_error(v) -> bool:
    return isinstance(v, str) and v.strip().startswith("#")


def _tint(fill: str) -> tuple[str, float] | None:
    m = re.fullmatch(r"(theme\d+):(-?\d+\.\d+)", fill or "")
    return (m.group(1), float(m.group(2))) if m else None


VALUE_HEADER_RE = re.compile(r"(embodied carbon|^ec )\s*-?.*kgco2e", re.I)


def _is_header_cell(v) -> bool:
    """True for a cell that makes its row a table header (unit text or 'Declared Units')."""
    if not isinstance(v, str):
        return False
    t = re.sub(r"\s+", " ", v).strip()
    return bool(VALUE_HEADER_RE.search(t)) or t.lower() == "declared units"


def _read_layout(row: SheetRow, title: SheetRow | None = None) -> Layout:
    """Work out which column holds what, from the text of one header row.

    `title` is the band-coloured row directly above the header, if there is
    one. It may carry labels (Comments, DQI ...) that the header leaves out.
    """
    sig = []
    lay = Layout(header_row=row.row, signature=())
    for source, is_main in ((row, True), (title, False)):
        if source is None:
            continue
        for idx in range(COL_G, N_COLS):
            raw = source.cells[idx]
            if not isinstance(raw, str) or not raw.strip():
                continue
            t = re.sub(r"\s+", " ", raw).strip()
            low = t.lower()
            if is_main:
                sig.append((_col_letter(idx), t))
            if "mass per declared unit" in low:
                lay.mass_col = lay.mass_col if lay.mass_col is not None else idx
            elif "declared unit" in low:
                lay.declared_col = lay.declared_col if lay.declared_col is not None else idx
            elif "of which carbon storage" in low:
                lay.storage_col = lay.storage_col if lay.storage_col is not None else idx
            elif VALUE_HEADER_RE.search(t):
                if lay.value_col is None and is_main:
                    lay.value_col = idx
                    lay.unit_text = t
                    lay.unit_kind = _unit_kind(low)
            elif low == "units":
                lay.unit_desc_col = lay.unit_desc_col if lay.unit_desc_col is not None else idx
            elif low.startswith("comments"):
                lay.comments_col = lay.comments_col if lay.comments_col is not None else idx
            elif low.startswith("dqi score"):
                lay.dqi_col = lay.dqi_col if lay.dqi_col is not None else idx
            elif low.startswith("dqi version"):
                lay.dqi_version_col = lay.dqi_version_col if lay.dqi_version_col is not None else idx
            elif low.startswith("name exactly as in ice db"):
                lay.name_col = lay.name_col if lay.name_col is not None else idx
            elif is_main:
                log.warning("row %d: header text not recognised in column %s: %r", row.row, _col_letter(idx), t)
    lay.signature = tuple(sig)
    e = row.cells[COL_E]
    if isinstance(e, str):
        m = re.search(r"kgco2e\s*per\s+(.+)$", re.sub(r"\s+", " ", e).strip(), re.I)
        if m:
            lay.table_declared_unit = m.group(1).strip()
    return lay


def _read_value_columns(row: SheetRow, lay: Layout, default_name_col: int | None):
    """Read a band row such as '% Cement Replacement - GGBS | 0.25 | 0.5 | 0.7'.

    Returns (list of ValueColumn, dqi_version_col). The DQI and name columns
    for each share are found by matching the share itself, not by position.
    """
    shares = [(idx, float(row.cells[idx])) for idx in (COL_G, COL_H, COL_I) if _is_number(row.cells[idx])]
    dqi_by_pct, dqi_version_col = {}, lay.dqi_version_col
    for idx in range(COL_I + 1, N_COLS):
        v = row.cells[idx]
        if not isinstance(v, str):
            continue
        t = re.sub(r"\s+", " ", v).strip().lower()
        m = re.match(r"dqi score \((\d+(?:\.\d+)?)%\)", t)
        if m:
            dqi_by_pct[round(float(m.group(1)), 6)] = idx
        elif t.startswith("dqi version"):
            dqi_version_col = idx
    name_by_share = {}
    if default_name_col is not None:
        for idx in range(default_name_col, N_COLS):
            v = row.cells[idx]
            if _is_number(v):
                name_by_share.setdefault(round(float(v), 6), idx)
    cols = [ValueColumn(idx, share, dqi_by_pct.get(round(share * 100, 6)), name_by_share.get(round(share, 6)))
            for idx, share in shares]
    return cols, dqi_version_col


def _unit_kind(header_text_lower: str) -> str | None:
    m = re.search(r"kgco2e\s*(?:/|per)\s*([a-z0-9]+)", header_text_lower)
    if not m:
        return None
    word = m.group(1)
    return {"kg": "kg", "tonne": "tonne", "t": "tonne", "m2": "m2", "unit": "unit"}.get(word)


def _looks_like_note(text: str) -> bool:
    low = text.lower()
    return (
        low.startswith(("note", "notes", "important note"))
        or "http" in low
        or "\n" in text
        or len(text) > 150
    )


def parse_summary(rows: list[SheetRow]) -> ParseResult:
    """Turn the rows of the summary sheet into factors, rejects and a structure log."""
    res = ParseResult()

    # which fill is "the header fill"? learn it from the header rows themselves
    header_rows = [r for r in rows if r.f == "Materials"]
    if not header_rows:
        raise IngestError("No header row found (no row with 'Materials' in column F).")
    header_fill = Counter(r.fill for r in header_rows).most_common(1)[0][0]
    header_tint = _tint(header_fill)
    theme = header_tint[0] if header_tint else None
    # 'Name Exactly as in ICE DB' sits in one fixed column; some headers leave the label out
    name_cols = Counter()
    for r in rows:
        for idx in range(COL_G, N_COLS):
            v = r.cells[idx]
            if isinstance(v, str) and v.strip().lower().startswith("name exactly as in ice db"):
                name_cols[idx] += 1
                break                                       # first labelled column of that row
    default_name_col = name_cols.most_common(1)[0][0] if name_cols else None

    def fill_rank(r: SheetRow) -> str | None:
        t = _tint(r.fill)
        if t and t[0] == theme:
            if 0.5 <= t[1] < 0.7:
                return "section"
            if 0.7 <= t[1] <= 0.9:
                return "band"
        return None

    def rank_of(r: SheetRow) -> str | None:
        """category / section / band / sub / None (= note) for a row with text in F only."""
        if r.size >= 20:
            return "category"
        by_fill = fill_rank(r)
        if by_fill:
            return by_fill
        t = _tint(r.fill)
        no_fill = (not r.fill) or (t is not None and t[0] != theme and t[1] == 0.0)   # white counts as none
        if r.bold and no_fill and not _looks_like_note(str(r.f)):
            return "sub"
        return None

    # ---- state while walking down the sheet
    cat = group = section = sub = band = None
    band_in_table = False
    colband = None                       # e.g. '% Cement Replacement - Blast Furnace Slag'
    value_columns: list[ValueColumn] | None = None
    multi_dqi_version_col = None
    last_text_band_row = -1
    version = None
    layout: Layout | None = None
    title_row: SheetRow | None = None    # band-coloured label row just above a header
    in_precast_model = False
    seen_category = False
    groups_seen: set = set()
    timber_including: list = []          # (key, value, storage, SheetRow, label)
    prev_row_number = 0

    def path() -> list[str]:
        return [p for p in (cat, group, section, sub, band, colband) if p]

    def reject(r: SheetRow, reason: str, detail: str = "", raw=None, name: str | None = None, extra_path=()):
        raw_v = r.cells[COL_G] if raw is None else raw
        res.rejected.append(Rejected(
            r.row, " > ".join(path() + list(extra_path)), name or _text(r.f) or "",
            "" if raw_v is None else str(raw_v)[:120], reason, detail))

    def note(r: SheetRow, kind: str):
        res.structure.append((r.row, kind, (_text(r.f) or _text(r.cells[COL_G]) or "")[:200]))

    def end_table():
        nonlocal layout, band, band_in_table, colband, value_columns, multi_dqi_version_col
        layout = None
        colband, value_columns, multi_dqi_version_col = None, None, None
        if band_in_table:
            band, band_in_table = None, False

    def accept(r: SheetRow, lay: Layout, raw, name_col, dqi_col, dqi_version_col, extra_path=()):
        """Check one value of one row and store it as a factor (or reject it)."""
        label = _text(r.f)
        shown = f"{label} [{extra_path[-1]}]" if extra_path else label
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            reject(r, R_BLANK, raw=raw, name=shown, extra_path=extra_path)
            return
        if not _is_number(raw):
            reject(r, R_TEXT, raw=raw, name=shown, extra_path=extra_path)
            return
        unit, declared_unit, mass, problem = _resolve_unit(r, lay)
        if problem:
            reject(r, R_UNIT, problem, raw=raw, name=shown, extra_path=extra_path)
            return
        value = float(raw) / 1000.0 if lay.unit_kind == "tonne" else float(raw)

        name_in_source = _text(r.cells[name_col]) if name_col is not None else None
        comments = _text(r.cells[lay.comments_col]) if lay.comments_col is not None else None
        dqi = r.cells[dqi_col] if dqi_col is not None else None
        dqi = float(dqi) if _is_number(dqi) and 0 <= dqi <= 1 else None
        dqi_version = _text(r.cells[dqi_version_col]) if dqi_version_col is not None else None
        if dqi_version and dqi_version.upper() == "N/A":
            dqi_version = None

        p = path() + list(extra_path)
        # ---- timber: one table without storage, one with
        if lay.storage_col is not None:
            heading = (section or "").lower()
            storage = r.cells[lay.storage_col]
            if "including carbon storage" in heading:
                key = TIMBER_SUFFIX_RE.sub("", label).lower()
                timber_including.append((key, float(raw), storage, r, label))
                return
            if "excluding carbon storage" not in heading:
                reject(r, R_TIMBER_NO_HEADING, raw=raw)
                return
            if _is_number(storage) and storage != 0:
                res.warnings.append(f"row {r.row}: 'excluding carbon storage' row has a storage value of {storage}")
            label = TIMBER_SUFFIX_RE.sub("", label)
            if name_in_source:
                name_in_source = TIMBER_SUFFIX_RE.sub("", name_in_source)
            p = [cat]                                   # the excluding/including heading is not part of the name

        if value < 0:
            reject(r, R_NEGATIVE, raw=raw, name=shown, extra_path=extra_path)
            return
        if unit == "kgCO2e/kg" and lay.declared_col is None:
            if name_in_source and SINGLE_ITEM_RE.search(name_in_source):
                unit, declared_unit = "kgCO2e/unit", name_in_source
                comments = f"{comments} {UNIT_FROM_NAME_NOTE}" if comments else UNIT_FROM_NAME_NOTE
                res.unit_from_name.append((r.row, label, name_in_source))
            elif SINGLE_ITEM_RE.search(label):
                reject(r, R_UNIT_CONFLICT, f"ICE DB name: {name_in_source!r}", raw=raw, name=shown, extra_path=extra_path)
                return
        res.factors.append(ParsedFactor(
            row=r.row, category=cat, subcategory_path=" > ".join(p), label=label,
            canonical_name=" > ".join(p + [label]), name_in_source=name_in_source,
            value=value, unit=unit, declared_unit=declared_unit, mass_per_item_kg=mass,
            source_value=float(raw), source_unit_text=lay.unit_text or "Declared Units",
            comments=comments, dqi_score=dqi, dqi_version=dqi_version, section_version=version,
        ))

    for r in rows:
        # a gap in row numbers is a blank row: it ends a table
        if r.row != prev_row_number + 1:
            if layout is not None:
                end_table()
            title_row = None
        prev_row_number = r.row

        f = r.f
        gi = r.cells[COL_G:COL_I + 1]
        f_is_text = isinstance(f, str) and f != ""
        only_f = f_is_text and all(v is None or (isinstance(v, str) and not v.strip()) for v in gi)

        # ---- the version the sheet says it is
        if f_is_text and SHEET_VERSION_LABEL_RE.match(f) and res.sheet_version is None:
            res.sheet_version = _text(r.cells[COL_G])
            note(r, "sheet version")
            continue

        # ---- the declared-unit drop-down for in-situ concrete
        if f_is_text and DECLARED_UNIT_SELECTOR_RE.match(f):
            res.declared_unit_selector = _text(r.cells[COL_G])
            note(r, "declared-unit selector")
            if (res.declared_unit_selector or "").lower() != "kg":
                raise IngestError(
                    f"Row {r.row}: the 'SELECT declared unit' cell reads {res.declared_unit_selector!r}, not 'kg'. "
                    "The in-situ concrete values in the sheet change with this setting, so they would be "
                    "loaded in the wrong unit. Open the file in Excel, set the cell to 'kg', save, and run again."
                )
            continue

        # ---- header row: a new table starts
        if f_is_text and (f == "Materials" or any(_is_header_cell(v) for v in gi)) and not in_precast_model:
            end_table()
            layout = _read_layout(r, title_row if (title_row is not None and title_row.row == r.row - 1) else None)
            if layout.name_col is None:
                layout.name_col = default_name_col
            res.header_variants[layout.signature] += 1
            note(r, "header")
            title_row = None
            continue

        # ---- rows with nothing in column F
        if not f_is_text:
            end_table()
            if in_precast_model and any(v is not None for v in r.cells[COL_G:COL_G + 4]):
                reject(r, R_PRECAST)
            elif seen_category and _is_number(r.cells[COL_G]):
                reject(r, R_NO_HEADER)
            else:
                note(r, "other")
            continue

        # ---- headings (text in F only)
        if only_f:
            is_version = f.lower().startswith("version:")
            rank = None if is_version else rank_of(r)
            if layout is not None and not is_version and rank not in ("category", "section"):
                if rank == "band":                          # sub-heading inside a table
                    band, band_in_table = f, True
                    colband, value_columns, multi_dqi_version_col = None, None, None
                    last_text_band_row = r.row
                    note(r, "band (in table)")
                elif _looks_like_note(f):                   # a sentence placed right under the last data row
                    note(r, "note")
                else:                                       # a material row with no value
                    reject(r, R_NO_VALUE_COL if layout.value_col is None and layout.declared_col is None else R_BLANK)
                continue
            end_table()
            if is_version:
                version = f.split(":", 1)[1].strip()
                note(r, "version")
            elif rank == "category":
                cat, group, section, sub, band = f, None, None, None, None
                version, in_precast_model, seen_category = None, False, True
                note(r, "category")
            elif not seen_category:
                note(r, "preamble")
            elif rank == "section":
                in_precast_model = False
                if f in GROUP_HEADINGS.get(cat, ()):
                    group, section, sub, band = f, None, None, None
                    groups_seen.add((cat, f))
                    note(r, "group")
                else:
                    section, sub, band = f, None, None
                    note(r, "section")
            elif in_precast_model:
                note(r, "note (precast model)")
            elif rank == "sub":
                if PRECAST_MODEL_RE.search(f):
                    in_precast_model = True
                    sub, band = None, None
                    note(r, "precast model start")
                else:
                    sub, band = f, None
                    note(r, "sub-section")
            elif rank == "band":
                band, band_in_table = f, False
                note(r, "band")
            else:
                note(r, "note")
            continue

        # ---- from here on the row has text in F and something in G..I
        if not seen_category:
            note(r, "preamble")
            continue
        if in_precast_model:
            reject(r, R_PRECAST)
            continue

        is_band_row = fill_rank(r) == "band"
        if layout is None:
            if any(_is_number(v) for v in gi) and not is_band_row:
                reject(r, R_NO_HEADER, "the table starts with a band row instead of a header, so no unit is given")
            else:
                if is_band_row:
                    title_row = r                           # may lend its labels to the header below
                note(r, "label")
            continue

        # ---- band row holding percentages: the table now has several value columns
        if is_band_row and any(_is_number(v) for v in gi):
            value_columns, multi_dqi_version_col = _read_value_columns(r, layout, default_name_col)
            if last_text_band_row != r.row - 1:
                band = None                                 # it replaces the band above; it does not sit inside it
            colband, band_in_table = f, True
            note(r, f"value columns: {', '.join(vc.label for vc in value_columns)}")
            continue

        if layout.value_col is None and layout.declared_col is None:
            reject(r, R_NO_VALUE_COL)
            continue

        # ---- a data row inside a table
        if value_columns:
            for vc in value_columns:
                accept(r, layout, r.cells[vc.col], vc.name_col, vc.dqi_col, multi_dqi_version_col,
                       extra_path=(vc.label,))
        else:
            value_col = layout.value_col if layout.value_col is not None else COL_H
            accept(r, layout, r.cells[value_col], layout.name_col, layout.dqi_col, layout.dqi_version_col)

    # ---- checks after the last row
    if res.declared_unit_selector is None:
        raise IngestError(
            "Could not find the 'SELECT declared unit' cell. The in-situ concrete values depend on it, "
            "so the unit of those values cannot be confirmed. Nothing was loaded."
        )
    for category, names in GROUP_HEADINGS.items():
        for name in names:
            if (category, name) not in groups_seen:
                res.warnings.append(f"group heading {name!r} was not found under {category!r}; paths there are one level shorter")

    _pair_timber(res, timber_including)
    _resolve_duplicates(res)
    return res


def _resolve_unit(r: SheetRow, lay: Layout):
    """Return (unit, declared_unit, mass_per_item_kg, problem). Units are never converted here."""
    if lay.declared_col is not None:
        du = r.cells[lay.declared_col]
        mass = r.cells[lay.mass_col] if lay.mass_col is not None else None
        if _is_number(du) and isinstance(mass, str):
            # some tables hold "1" under Declared Units and "m2" under Mass: together they read "1 m2"
            if du != 1:
                return None, None, None, f"declared unit is {du} x {mass!r}; only 1 x unit is understood"
            du, mass = mass, None
        du_text = (_text(du) or "").lower()
        if not du_text or _is_excel_error(du_text):
            return None, None, None, "declared unit is missing"
        mass = float(mass) if _is_number(mass) and mass > 0 else None
        if du_text == "kg":
            if mass is not None and abs(mass - 1.0) > 1e-9:
                return None, None, None, f"declared unit is kg but mass per declared unit is {mass}"
            return "kgCO2e/kg", None, None, None
        if du_text == "m2":
            return "kgCO2e/m2", None, mass, None
        return "kgCO2e/unit", du_text, mass, None
    if lay.unit_kind in ("kg", "tonne"):
        return "kgCO2e/kg", None, None, None
    if lay.unit_kind == "m2":
        return "kgCO2e/m2", None, None, None
    if lay.unit_kind == "unit":
        desc = r.cells[lay.unit_desc_col] if lay.unit_desc_col is not None else None
        desc = None if (desc is None or _is_excel_error(desc)) else _text(desc)
        desc = desc or lay.table_declared_unit
        if not desc:
            return None, None, None, "per-unit value with no description of the unit"
        return "kgCO2e/unit", desc, None, None
    return None, None, None, f"header unit not recognised: {lay.unit_text!r}"


def _pair_timber(res: ParseResult, including: list) -> None:
    """Attach carbon storage (from the 'including' table) to the matching 'excluding' material."""
    by_key = {}
    for fct in res.factors:
        if fct.category and fct.subcategory_path == fct.category:
            by_key.setdefault(fct.label.lower(), fct)
    for key, incl_value, storage, r, label in including:
        target = by_key.get(key)
        if target is None:
            res.rejected.append(Rejected(r.row, "Timber > Data including carbon storage", label, str(incl_value), R_TIMBER_UNPAIRED))
            continue
        if not _is_number(storage):
            res.warnings.append(f"row {r.row}: {label!r} has no carbon storage value; no biogenic figure stored")
            continue
        # cross-check the sheet against itself: including = excluding + storage
        expected = target.source_value + float(storage)
        if abs(expected - incl_value) > 1e-6 * max(1.0, abs(incl_value)):
            res.warnings.append(
                f"row {r.row}: {label!r}: including ({incl_value}) is not excluding ({target.source_value}) "
                f"+ storage ({storage}); biogenic figure NOT stored"
            )
            continue
        if storage > 0:
            res.warnings.append(f"row {r.row}: {label!r} has a positive storage value ({storage}); not stored")
            continue
        target.biogenic_kgco2e_per_kg = float(storage)
        target.biogenic_row = r.row
        res.timber_pairs += 1


def _squash(text: str) -> str:
    return re.sub(r"[\s,]+", "", (text or "").lower())


def _resolve_duplicates(res: ParseResult) -> None:
    """Make every material unique, on purpose and visibly. Three rules, in this order.

    Rule A  Same 'Name Exactly as in ICE DB', same unit, same value, on rows that sit side by
            side under the same heading but have DIFFERENT labels. One of the labels is wrong
            in the source. Keep the row whose label appears in the ICE DB name; reject the other.
    Rule B  Same ICE DB name, same unit, same value anywhere else in the sheet. The sheet shows
            one datapoint in two places (e.g. the "0% replacement" column repeats the CEM I
            table). Keep the first, reject the rest.
    Rule C  Same built name but different values. They are different materials whose row label
            alone does not say so ("As above but ..."). Add the ICE DB name in square brackets;
            if that is missing or still clashes, add the sheet row number.
    A shared ICE DB name with DIFFERENT values is never merged; it is reported as a warning.
    """
    drop: dict[int, None] = {}

    def same_value(a: ParsedFactor, b: ParsedFactor) -> bool:
        return a.unit == b.unit and abs(a.value - b.value) <= 1e-12 * max(1.0, abs(b.value))

    def reject(m: ParsedFactor, reason: str, kept: ParsedFactor, text: str):
        drop[id(m)] = None
        res.rejected.append(Rejected(m.row, m.subcategory_path, m.label, str(m.source_value), reason,
                                     f"kept row {kept.row}: {kept.canonical_name}"))
        res.duplicates.append({"name": m.name_in_source or m.canonical_name, "rows": [kept.row, m.row], "resolution": text})

    # ---- Rules A and B: by ICE DB name
    by_name: dict[tuple, list[ParsedFactor]] = {}
    for fct in res.factors:
        if fct.name_in_source:
            by_name.setdefault((fct.name_in_source.lower(), fct.unit), []).append(fct)
    for members in by_name.values():
        if len(members) < 2:
            continue
        if not all(same_value(m, members[0]) for m in members[1:]):
            res.warnings.append(
                f"ICE DB name {members[0].name_in_source!r} is used on rows {[m.row for m in members]} "
                "with different values; all were kept as separate materials")
            continue
        alive = list(members)
        fits = [m for m in alive if _squash(m.label) in _squash(m.name_in_source)]
        for m in list(alive):
            siblings = [o for o in alive if o is not m and o.subcategory_path == m.subcategory_path and o.label != m.label]
            if siblings and m not in fits and any(o in fits for o in siblings):
                kept = next(o for o in siblings if o in fits)
                reject(m, R_LABEL_MISMATCH, kept,
                       f"Rule A: label {m.label!r} is not in the ICE DB name and the value equals row {kept.row}; rejected")
                alive.remove(m)
        first = alive[0]
        for m in alive[1:]:
            reject(m, R_DUPLICATE, first, f"Rule B: same datapoint shown again under '{m.subcategory_path}'; kept row {first.row}")
    res.factors = [f for f in res.factors if id(f) not in drop]

    # ---- Rule C: by built name
    groups: dict[str, list[ParsedFactor]] = {}
    for fct in res.factors:
        groups.setdefault(fct.canonical_name.lower(), []).append(fct)
    for members in groups.values():
        if len(members) < 2:
            continue
        old = members[0].canonical_name
        distinct = [members[0]]
        for m in members[1:]:
            twin = next((d for d in distinct if same_value(m, d) and (m.name_in_source or "") == (d.name_in_source or "")), None)
            if twin is not None:
                reject(m, R_DUPLICATE, twin, f"Rule B: exact repeat of row {twin.row}")
            else:
                distinct.append(m)
        if len(distinct) > 1:
            for m in distinct:
                extra = m.name_in_source if m.name_in_source and m.name_in_source.lower() != m.label.lower() else None
                if extra:
                    m.canonical_name = f"{m.canonical_name} [{extra}]"
            names = Counter(m.canonical_name.lower() for m in distinct)
            for m in distinct:
                if names[m.canonical_name.lower()] > 1:
                    m.canonical_name = f"{m.canonical_name} [row {m.row}]"
            res.duplicates.append({"name": old, "rows": [m.row for m in distinct],
                                   "resolution": "Rule C: same label, different values; ICE DB name added in [ ] to tell them apart"})
    res.factors = [f for f in res.factors if id(f) not in drop]
    clash = [n for n, c in Counter(f.canonical_name.lower() for f in res.factors).items() if c > 1]
    if clash:
        raise IngestError(f"Duplicate material names remain after resolution: {clash[:5]}")


# ---------------------------------------------------------------- staging files
def write_staging(res: ParseResult, staging_dir: Path) -> None:
    staging_dir.mkdir(parents=True, exist_ok=True)
    with open(staging_dir / "ice_parsed.csv", "w", newline="", encoding="utf-8") as fh:
        cols = list(ParsedFactor.__dataclass_fields__)
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for f in res.factors:
            w.writerow(asdict(f))
    with open(staging_dir / "ice_rejected.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["row", "path", "name", "raw_value", "reason", "detail"])
        for r in sorted(res.rejected, key=lambda x: x.row):
            w.writerow([r.row, r.path, r.name, r.raw_value, r.reason, r.detail])
    with open(staging_dir / "ice_structure.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["row", "kind", "text"])
        w.writerows(res.structure)


# ---------------------------------------------------------------- step 3: load
def load_to_db(conn, res: ParseResult) -> dict:
    """Write the parsed factors to the database, all or nothing."""
    if EXPECTED_SHEET_VERSION.lower() not in (res.sheet_version or "").lower():
        raise IngestError(
            f"This script loads ICE Educational {EXPECTED_SHEET_VERSION}, but the sheet's 'ICE DB Version:' cell reads "
            f"{res.sheet_version!r}. Nothing was loaded: the factors would be stored under the wrong version and date."
        )
    exists = conn.execute(
        "SELECT id FROM data_sources WHERE name = ? AND version = ?", (ICE_SOURCE["name"], ICE_SOURCE["version"])
    ).fetchone()
    if exists:
        raise IngestError(
            f"{ICE_SOURCE['name']} {ICE_SOURCE['version']} is already in this database. Factors are never "
            "deleted or overwritten. To reload, delete the database file in db/ and run the ingestion again."
        )
    counts = Counter()
    with db.transaction(conn):
        src_id = conn.execute(
            "INSERT INTO data_sources (name, publisher, version, licence, attribution_text, url, redistributable) "
            "VALUES (:name, :publisher, :version, :licence, :attribution_text, :url, :redistributable)", ICE_SOURCE
        ).lastrowid
        for f in res.factors:
            mat_id = conn.execute(
                "INSERT INTO materials (canonical_name, category, subcategory_path, name_in_source, data_source_id, mass_per_item_kg) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (f.canonical_name, f.category, f.subcategory_path, f.name_in_source, src_id, f.mass_per_item_kg),
            ).lastrowid
            counts["materials"] += 1
            conn.execute(
                "INSERT INTO carbon_factors (material_id, value, unit, declared_unit, lifecycle_module, data_source_id, "
                "source_section_version, dqi_score, dqi_version, comments, valid_from) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (mat_id, f.value, f.unit, f.declared_unit, LIFECYCLE_MODULE, src_id, f.section_version,
                 f.dqi_score, f.dqi_version, f.comments, ICE_VALID_FROM),
            )
            counts["carbon_factors"] += 1
            if f.biogenic_kgco2e_per_kg is not None:
                conn.execute(
                    "INSERT INTO biogenic_carbon (material_id, value_kgco2e_per_kg, data_source_id, comments, valid_from) "
                    "VALUES (?, ?, ?, ?, ?)",
                    (mat_id, f.biogenic_kgco2e_per_kg, src_id,
                     "ICE 'Of which carbon storage' column, from the 'including carbon storage' table", ICE_VALID_FROM),
                )
                counts["biogenic_carbon"] += 1
    return dict(counts)


# ---------------------------------------------------------------- report
def summary_text(rows, res: ParseResult, loaded: dict | None, density_stats: dict | None = None) -> str:
    out = []
    out.append("=" * 72)
    out.append("ICE ingestion summary")
    out.append("=" * 72)
    out.append(f"Non-empty rows read from '{SUMMARY_SHEET}':   {len(rows)}")
    out.append(f"Header rows (tables) found:               {sum(res.header_variants.values())}  "
               f"in {len(res.header_variants)} different layouts")
    out.append(f"'ICE DB Version:' cell reads:             {res.sheet_version!r}  (must contain '{EXPECTED_SHEET_VERSION}')")
    out.append(f"'SELECT declared unit' cell reads:        {res.declared_unit_selector!r}  (must be 'kg')")
    out.append("")
    sheet_rows = len({f.row for f in res.factors})
    out.append(f"Factors accepted per category  ({len(res.factors)} factors from {sheet_rows} sheet rows; "
               "a row can hold up to 3 values):")
    per_cat = Counter(f.category for f in res.factors)
    for c in sorted(per_cat):
        out.append(f"   {c:<42s}{per_cat[c]:>5d}")
    out.append(f"   {'TOTAL':<42s}{len(res.factors):>5d}")
    by_unit = Counter(f.unit for f in res.factors)
    out.append("   by unit: " + ", ".join(f"{u} = {n}" for u, n in sorted(by_unit.items())))
    tonne = sum(1 for f in res.factors if "tonne" in f.source_unit_text.lower())
    for row, label, name in res.unit_from_name:
        out.append(f"   stored per unit because the ICE DB name says one item (table header says per kg): "
                   f"row {row} {label!r} -> declared unit {name!r}")
    out.append(f"   converted from per tonne to per kg (divided by 1000): {tonne}")
    v_old = Counter(f.section_version for f in res.factors if f.section_version)
    if v_old:
        out.append("   from sections carrying an older version label: "
                   + ", ".join(f"{v} = {n}" for v, n in sorted(v_old.items())))
    out.append("")
    out.append(f"Rows rejected: {len(res.rejected)}   (listed in data/staging/ice_rejected.csv)")
    for reason, n in Counter(r.reason for r in res.rejected).most_common():
        out.append(f"   {n:>4d}  {reason}")
    out.append("")
    out.append(f"Timber: {res.timber_pairs} 'including carbon storage' rows paired with an 'excluding' material "
               "(storage kept separately as biogenic carbon)")
    out.append("")
    out.append(f"Duplicates found and resolved: {len(res.duplicates)}   (every one is listed in ice_rejected.csv or below)")
    by_rule = Counter(d["resolution"].split(":")[0] for d in res.duplicates)
    explain = {
        "Rule A": "label disagrees with the ICE DB name and repeats a neighbour's value -> rejected",
        "Rule B": "same ICE DB name and same value shown twice in the sheet -> first kept",
        "Rule C": "same built name, different values -> ICE DB name added to keep both",
    }
    for rule in sorted(by_rule):
        out.append(f"   {by_rule[rule]:>4d}  {rule}: {explain.get(rule, '')}")
    for d in res.duplicates:
        if not d["resolution"].startswith("Rule B"):
            out.append(f"         rows {d['rows']}: {str(d['name'])[:90]}")
    if res.warnings:
        out.append("")
        out.append(f"Warnings: {len(res.warnings)}")
        for w in res.warnings:
            out.append(f"   {w}")
    if density_stats is not None:
        out.append("")
        out.append(f"Density candidates extracted: {density_stats['total']}  from {density_stats['sheets']} profile sheets "
                   "(data/staging/density_candidates.csv)")
        out.append(f"   outside the sanity range:     {density_stats['flagged']}")
        out.append(f"   inside the sanity range:      {density_stats['ok']}")
        out.append(f"   sheet has no range defined:   {density_stats['no_range']}")
    out.append("")
    if loaded is None:
        out.append("Rows loaded: 0 (--no-load was given)")
    else:
        out.append("Rows loaded into the database:")
        for t in ("materials", "carbon_factors", "biogenic_carbon"):
            out.append(f"   {t:<18s}{loaded.get(t, 0):>5d}")
        out.append(f"   every factor: lifecycle module {LIFECYCLE_MODULE}, valid_from {ICE_VALID_FROM}")
    out.append("=" * 72)
    return "\n".join(out)


# ---------------------------------------------------------------- command line
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Ingest the ICE Database summary sheet.")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--db", type=Path, default=db.DEFAULT_DB_PATH)
    ap.add_argument("--staging", type=Path, default=STAGING_DIR)
    ap.add_argument("--no-load", action="store_true", help="parse and write staging files only")
    ap.add_argument("--no-densities", action="store_true", help="skip the density tables in the profile sheets")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")

    try:
        log.info("Step 1/5  reading sheet '%s' from %s", SUMMARY_SHEET, args.xlsx.name)
        rows = read_summary_sheet(args.xlsx)
        write_raw_csv(rows, args.staging / "ice_summary_raw.csv")
        log.info("          %d non-empty rows written to %s", len(rows), args.staging / "ice_summary_raw.csv")

        log.info("Step 2/5  parsing")
        res = parse_summary(rows)
        log.info("          %d header rows in %d layouts; %d factors accepted; %d rows rejected; %d notes/headings",
                 sum(res.header_variants.values()), len(res.header_variants), len(res.factors),
                 len(res.rejected), len(res.structure))

        log.info("Step 3/5  writing staging files")
        write_staging(res, args.staging)
        log.info("          ice_parsed.csv (%d rows), ice_rejected.csv (%d rows), ice_structure.csv (%d rows)",
                 len(res.factors), len(res.rejected), len(res.structure))

        density_stats = None
        if not args.no_densities:
            from carbon_cost.ingest import densities
            log.info("Step 4/5  extracting density candidates from the material profile sheets")
            cands = densities.extract_density_candidates(args.xlsx)
            density_stats = densities.write_candidates(cands, args.staging / "density_candidates.csv")
            log.info("          %d candidates from %d sheets; %d outside their sanity range",
                     density_stats["total"], density_stats["sheets"], density_stats["flagged"])
        else:
            log.info("Step 4/5  skipped (--no-densities)")

        loaded = None
        if not args.no_load:
            log.info("Step 5/5  loading into %s", args.db)
            conn = db.connect(args.db) if Path(args.db).exists() else db.init_db(args.db)
            try:
                loaded = load_to_db(conn, res)
            finally:
                conn.close()
            log.info("          loaded %s", loaded)
        else:
            log.info("Step 5/5  skipped (--no-load)")
    except IngestError as exc:
        log.error("INGESTION STOPPED: %s", exc)
        return 1

    print()
    print(summary_text(rows, res, loaded, density_stats))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
